import os
from docx import Document


class TemplateService:
    """
    Populates a .docx template by replacing placeholder tokens
    (e.g. {{Q1_ANSWER}}) with the actual RAG answers.

    Why run merging?
    Word often splits a placeholder like {{Q1_ANSWER}} across multiple
    XML runs when the user types or edits it. A naive paragraph.text
    replace silently fails in those cases.  We collapse all runs in each
    paragraph to a single string, do the replacement, then write back.
    Table cells are handled separately — the original code missed these.
    """

    def populate(self, template_path: str, output_path: str, answers: dict[str, str]) -> None:
        """
        Replace all placeholder tokens in the template and save to output_path.

        Parameters
        ----------
        template_path : path to the template .docx (contains {{Q1_ANSWER}} etc.)
        output_path   : where to save the populated document
        answers       : { "{{Q1_ANSWER}}": "synthesised answer text", ... }
        """
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

        doc = Document(template_path)

        # Body paragraphs
        for para in doc.paragraphs:
            self._replace_in_paragraph(para, answers)

        # Table cells (header rows, answer boxes built as tables)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        self._replace_in_paragraph(para, answers)

        doc.save(output_path)

    # ── Private ───────────────────────────────────────────────────────────────

    def _replace_in_paragraph(self, paragraph, answers: dict[str, str]) -> None:
        """
        Collapse runs → replace placeholders → write back into run[0].
        Clears extra runs so no duplicate text appears.
        """
        full_text = "".join(run.text for run in paragraph.runs)

        replaced = False
        for placeholder, value in answers.items():
            if placeholder in full_text:
                full_text = full_text.replace(placeholder, value)
                replaced = True

        if replaced and paragraph.runs:
            paragraph.runs[0].text = full_text
            for run in paragraph.runs[1:]:
                run.text = ""
